import os
import re
from datetime import datetime
from zoneinfo import ZoneInfo
from io import BytesIO

import requests
from openpyxl import load_workbook
import smtplib
from email.mime.text import MIMEText


# =========================
# Settings / Secrets
# =========================
EXCEL_URL = os.environ.get("EXCEL_URL", "").strip()

SMTP_HOST = os.environ.get("SMTP_HOST", "").strip()
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))
SMTP_USER = os.environ.get("SMTP_USER", "").strip()
SMTP_PASS = os.environ.get("SMTP_PASS", "").strip()
MAIL_FROM = os.environ.get("MAIL_FROM", "").strip()
MAIL_TO = os.environ.get("MAIL_TO", "").strip()

PAGES_BASE_URL = os.environ.get("PAGES_BASE_URL", "").strip()  # optional
TZ = ZoneInfo("Asia/Muscat")
AUTO_OPEN_ACTIVE_SHIFT_IN_FULL = True
# Excel sheets
DEPARTMENTS = [
    ("Officers", "Officers"),
    ("Supervisors", "Supervisors"),
    ("Load Control", "Load Control"),
    ("Export Checker", "Export Checker"),
    ("Export Operators", "Export Operators"),
]

# For day-row matching only
DAYS = ["SUN", "MON", "TUE", "WED", "THU", "FRI", "SAT"]

SHIFT_MAP = {
    "MN06": ("🌅 Morning (MN06)", "صباح"),
    "ME06": ("🌅 Morning (ME06)", "صباح"),
    "ME07": ("🌅 Morning (ME07)", "صباح"),
    "MN12": ("🌆 Afternoon (MN12)", "ظهر"),
    "AN13": ("🌆 Afternoon (AN13)", "ظهر"),
    "AE14": ("🌆 Afternoon (AE14)", "ظهر"),
    "NN21": ("🌙 Night (NN21)", "ليل"),
    "NE22": ("🌙 Night (NE22)", "ليل"),
}

GROUP_ORDER = ["صباح", "ظهر", "ليل", "مناوبات", "راحة", "إجازات", "تدريب", "أخرى"]


# =========================
# Helpers
# =========================
def clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\u00A0", " ")).strip()

def to_western_digits(s: str) -> str:
    if s is None:
        return ""
    s = str(s)
    arabic = {'٠':'0','١':'1','٢':'2','٣':'3','٤':'4','٥':'5','٦':'6','٧':'7','٨':'8','٩':'9'}
    farsi  = {'۰':'0','۱':'1','۲':'2','۳':'3','۴':'4','۵':'5','۶':'6','۷':'7','۸':'8','۹':'9'}
    mp = {**arabic, **farsi}
    return "".join(mp.get(ch, ch) for ch in s)

def norm(s) -> str:
    return clean(to_western_digits(s))

def looks_like_time(s: str) -> bool:
    up = norm(s).upper()
    return bool(
        re.match(r"^\d{3,4}\s*H?\s*-\s*\d{3,4}\s*H?$", up)
        or re.match(r"^\d{3,4}\s*H$", up)
        or re.match(r"^\d{3,4}$", up)
    )

def looks_like_employee_name(s: str) -> bool:
    v = norm(s)
    if not v:
        return False
    up = v.upper()
    if looks_like_time(up):
        return False
    if re.search(r"(ANNUAL\s*LEAVE|SICK\s*LEAVE|REST\/OFF\s*DAY|REST|OFF\s*DAY|TRAINING|STANDBY)", up):
        return False
    # قوي: اسم - رقم
    if re.search(r"-\s*\d{3,}", v) and re.search(r"[A-Za-z\u0600-\u06FF]", v):
        return True
    # بديل: كلمتين أو أكثر
    parts = [p for p in v.split(" ") if p]
    return bool(re.search(r"[A-Za-z\u0600-\u06FF]", v) and len(parts) >= 2)

def looks_like_shift_code(s: str) -> bool:
    v = norm(s).upper()
    if not v:
        return False
    if looks_like_time(v):
        return False
    if v in ["OFF", "O", "LV", "TR", "ST", "SL", "AL", "STM", "STN"]:
        return True
    if re.match(r"^(MN|AN|NN|NT|ME|AE|NE)\d{1,2}", v):
        return True
    if re.search(r"(ANNUAL\s*LEAVE|SICK\s*LEAVE|REST\/OFF\s*DAY|REST|OFF\s*DAY|TRAINING|STANDBY)", v):
        return True
    return False

def map_shift(code: str):
    c0 = norm(code)
    c = c0.upper()
    if not c or c == "0":
        return ("-", "أخرى")

    if c == "AL" or "ANNUAL LEAVE" in c:
        return ("🏖️ Leave", "إجازات")
    if c == "SL" or "SICK LEAVE" in c:
        return ("🤒 Sick Leave", "إجازات")
    if c == "LV":
        return ("🏖️ Leave", "إجازات")
    if c in ["TR"] or "TRAINING" in c:
        return ("📚 Training", "تدريب")
    if c in ["ST", "STM", "STN"] or "STANDBY" in c:
        return ("🧍 Standby", "مناوبات")
    if c in ["OFF", "O"] or re.search(r"(REST|OFF\s*DAY|REST\/OFF)", c):
        return ("🛌 Off Day", "راحة")

    if c in SHIFT_MAP:
        return SHIFT_MAP[c]

    return (c0, "أخرى")

def current_shift_key(now: datetime) -> str:
    # 21:00–04:59 Night, 14:00–20:59 Afternoon, else Morning
    t = now.hour * 60 + now.minute
    if t >= 21 * 60 or t < 5 * 60:
        return "ليل"
    if t >= 14 * 60:
        return "ظهر"
    return "صباح"

def download_excel(url: str) -> bytes:
    r = requests.get(url, timeout=60)
    r.raise_for_status()
    return r.content

def infer_pages_base_url():
    return "https://khalidsaif912.github.io/roster-site"


# =========================
# Detect rows/cols (Days row + Date numbers row)
# =========================
def _row_values(ws, r: int):
    return [norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]

def _count_day_tokens(vals) -> int:
    ups = [v.upper() for v in vals if v]
    count = 0
    for d in DAYS:
        if any(d in x for x in ups):
            count += 1
    return count

def _is_date_number(v: str) -> bool:
    v = norm(v)
    if not v:
        return False
    if re.match(r"^\d{1,2}(\.0)?$", v):
        n = int(float(v))
        return 1 <= n <= 31
    return False

def find_days_and_dates_rows(ws, scan_rows: int = 80):
    """
    يبحث عن صف فيه SUN..SAT بكثرة ثم صف تحته فيه أرقام 1..31
    """
    max_r = min(ws.max_row, scan_rows)
    days_row = None

    for r in range(1, max_r + 1):
        vals = _row_values(ws, r)
        if _count_day_tokens(vals) >= 3:
            days_row = r
            break

    if not days_row:
        return None, None

    date_row = None
    for r in range(days_row + 1, min(days_row + 4, ws.max_row) + 1):
        vals = _row_values(ws, r)
        nums = sum(1 for v in vals if _is_date_number(v))
        if nums >= 5:
            date_row = r
            break

    return days_row, date_row

def find_day_col(ws, days_row: int, date_row: int, today_dow: int, today_day: int):
    """
    يثبت العمود الصحيح باستخدام اليوم + رقم التاريخ
    """
    if not days_row or not date_row:
        return None

    day_key = DAYS[today_dow]
    # Prefer (day + date) match
    for c in range(1, ws.max_column + 1):
        top = norm(ws.cell(row=days_row, column=c).value).upper()
        bot = norm(ws.cell(row=date_row, column=c).value)
        if day_key in top and _is_date_number(bot) and int(float(bot)) == today_day:
            return c

    # Fallback: date-only
    for c in range(1, ws.max_column + 1):
        bot = norm(ws.cell(row=date_row, column=c).value)
        if _is_date_number(bot) and int(float(bot)) == today_day:
            return c

    return None

def find_employee_col(ws, start_row: int, max_scan_rows: int = 200):
    scores = {}
    r_end = min(ws.max_row, start_row + max_scan_rows)
    for r in range(start_row, r_end + 1):
        for c in range(1, ws.max_column + 1):
            if looks_like_employee_name(ws.cell(row=r, column=c).value):
                scores[c] = scores.get(c, 0) + 1
    if not scores:
        return None
    return max(scores.items(), key=lambda kv: kv[1])[0]


# =========================
# EXACT DESIGN (as you provided)
# =========================
CSS = r"""
    /* ═══════ RESET ═══════ */
    body {
      margin:0; padding:0;
      background:#eef1f7;
      font-family:'Segoe UI', system-ui, -apple-system, BlinkMacSystemFont, Roboto, Helvetica, Arial, sans-serif;
      color:#0f172a;
      -webkit-font-smoothing:antialiased;
    }
    * { box-sizing:border-box; }

    /* ═══════ WRAP ═══════ */
    .wrap { max-width:680px; margin:0 auto; padding:16px 14px 28px; }

    /* ═══════ HEADER ═══════ */
    .header {
      background:linear-gradient(135deg, #1e40af 0%, #1976d2 50%, #0ea5e9 100%);
      color:#fff;
      padding:26px 18px 24px;
      border-radius:20px;
      text-align:center;
      box-shadow:0 8px 28px rgba(30,64,175,.25);
      position:relative;
      overflow:hidden;
    }
    .header::before {
      content:''; position:absolute;
      top:-30px; right:-40px;
      width:140px; height:140px;
      border-radius:50%;
      background:rgba(255,255,255,.08);
    }
    .header::after {
      content:''; position:absolute;
      bottom:-50px; left:-30px;
      width:160px; height:160px;
      border-radius:50%;
      background:rgba(255,255,255,.06);
    }
    .header h1 { margin:0; font-size:24px; font-weight:800; position:relative; z-index:1; letter-spacing:-.3px; }
    .header .dateTag {
      display:inline-block; margin-top:10px;
      background:rgba(255,255,255,.18);
      padding:5px 18px; border-radius:30px;
      font-size:13px; font-weight:600; letter-spacing:.3px;
      position:relative; z-index:1;
    }

    /* ═══════ SUMMARY BAR ═══════ */
    .summaryBar { display:flex; justify-content:center; gap:12px; margin-top:14px; }
    .summaryChip {
      background:#fff;
      border:1px solid rgba(15,23,42,.1);
      border-radius:14px;
      padding:10px 20px;
      text-align:center;
      box-shadow:0 2px 8px rgba(15,23,42,.06);
    }
    .summaryChip .chipVal { font-size:22px; font-weight:900; color:#1e40af; }
    .summaryChip .chipLabel { font-size:11px; font-weight:600; color:#64748b; text-transform:uppercase; letter-spacing:.6px; margin-top:2px; }

    /* ═══════ DEPARTMENT CARD ═══════ */
    .deptCard {
      margin-top:18px;
      background:#fff;
      border-radius:18px;
      overflow:hidden;
      border:1px solid rgba(15,23,42,.07);
      box-shadow:0 4px 18px rgba(15,23,42,.08);
    }
    .deptHead {
      display:flex;
      align-items:center;
      gap:12px;
      padding:14px 16px;
      background:#fff;
    }
    .deptIcon {
      width:40px; height:40px;
      border-radius:12px;
      display:flex; align-items:center; justify-content:center;
      flex-shrink:0;
    }
    .deptTitle { font-size:18px; font-weight:800; color:#1e293b; flex:1; letter-spacing:-.2px; }
    .deptBadge { min-width:48px; padding:6px 10px; border-radius:12px; text-align:center; }

    /* ═══════ SHIFT STACK ═══════ */
    .shiftStack { padding:10px; display:flex; flex-direction:column; gap:8px; }

    /* ═══════ SHIFT CARD — <details> ═══════ */
    .shiftCard {
      border-radius:14px;
      overflow:hidden;
    }

    .shiftSummary {
      display:flex;
      align-items:center;
      gap:10px;
      padding:11px 14px;
      cursor:pointer;
      list-style:none;
      -webkit-appearance:none;
      appearance:none;
      user-select:none;
    }
    .shiftSummary::-webkit-details-marker { display:none; }
    .shiftSummary::marker              { display:none; }

    .shiftIcon  { font-size:20px; line-height:1; flex-shrink:0; }
    .shiftLabel { font-size:15px; font-weight:800; flex:1; letter-spacing:-.1px; }
    .shiftCount {
      font-size:13px; font-weight:800;
      padding:3px 10px; border-radius:20px;
      flex-shrink:0;
    }

    /* chevron يدور لما يفتح */
    .shiftSummary::after {
      content:'▾';
      font-size:14px;
      color:#94a3b8;
      transition:transform .2s;
      flex-shrink:0;
    }
    .shiftCard[open] .shiftSummary::after {
      transform:rotate(180deg);
    }

    .shiftBody { background:rgba(255,255,255,.7); }

    /* ── employee row ── */
    .empRow {
      display:flex;
      align-items:center;
      justify-content:space-between;
      padding:9px 16px;
      border-top:1px solid rgba(15,23,42,.06);
    }
    .empRowAlt { background:rgba(15,23,42,.02); }
    .empName  { font-size:15px; font-weight:700; color:#1e293b; }
    .empStatus { font-size:13px; font-weight:600; }

    /* ═══════ CTA ═══════ */
    .btnWrap { margin-top:20px; text-align:center; }
    .btn {
      display:inline-block;
      padding:14px 38px;
      border-radius:16px;
      background:linear-gradient(135deg, #1e40af, #1976d2);
      color:#fff !important;
      text-decoration:none;
      font-weight:800;
      font-size:15px;
      box-shadow:0 6px 20px rgba(30,64,175,.3);
    }

    /* ═══════ FOOTER ═══════ */
    .footer { margin-top:18px; text-align:center; font-size:12px; color:#94a3b8; padding:12px 0; line-height:1.9; }
    .footer strong { color:#64748b; }

    /* ═══════ MOBILE ═══════ */
    @media (max-width:480px){
      .wrap            { padding:12px 10px 22px; }
      .header h1       { font-size:21px; }
      .deptTitle       { font-size:16px; }
      .empName         { font-size:14px; }
      .empStatus       { font-size:12px; }
      .shiftLabel      { font-size:14px; }
      .summaryBar      { gap:8px; }
      .summaryChip     { padding:8px 14px; }
      .summaryChip .chipVal { font-size:19px; }
    }
"""

DEPT_COLORS = ["#2563eb", "#7c3aed", "#0891b2", "#059669", "#dc2626", "#ea580c"]


# Email colors per department (to match site)
DEPT_EMAIL_COLORS = {
    "Officers": "#2563eb",
    "Supervisors": "#7c3aed",
    "Load Control": "#0891b2",
    "Export Checker": "#059669",
    "Export Operators": "#dc2626",
}

SVG_ICON = """
<svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round">
  <path d="M3 21h18M3 10h18M5 21V10l7-6 7 6v11"/>
  <rect x="9" y="14" width="2" height="3"/>
  <rect x="13" y="14" width="2" height="3"/>
</svg>
"""

def shift_style(grp: str, label_text: str):
    """
    Returns: (shift_title, icon, border_color, bg_color, text_color, count_bg)
    """
    if grp == "صباح":
        return ("Morning", "☀️", "#f59e0b44", "#fef3c7", "#92400e", "#f59e0b22")
    if grp == "ظهر":
        return ("Afternoon", "🌤️", "#f9731644", "#ffedd5", "#9a3412", "#f9731622")
    if grp == "ليل":
        return ("Night", "🌙", "#8b5cf644", "#ede9fe", "#5b21b6", "#8b5cf622")
    if grp == "راحة":
        return ("Off Day", "🛋️", "#6366f144", "#e0e7ff", "#3730a3", "#6366f122")
    if grp == "إجازات":
        # differentiate sick via label
        if "SICK" in label_text.upper() or "🤒" in label_text:
            return ("Sick Leave", "🏥", "#ef444444", "#fee2e2", "#991b1b", "#ef444422")
        return ("Annual Leave", "✈️", "#10b98144", "#d1fae5", "#065f46", "#10b98122")
    if grp == "تدريب":
        return ("Training", "📚", "#0ea5e944", "#e0f2fe", "#075985", "#0ea5e922")
    if grp == "مناوبات":
        return ("Standby", "🧍", "#94a3b844", "#f1f5f9", "#334155", "#94a3b822")
    return ("Other", "📌", "#64748b44", "#f8fafc", "#334155", "#64748b22")

def dept_card_html(dept_name: str, dept_color: str, buckets: dict, open_group: str | None = None):
    total = sum(len(buckets.get(g, [])) for g in GROUP_ORDER)
    shift_blocks = []

    for g in GROUP_ORDER:
        rows = buckets.get(g, [])
        if not rows:
            continue

        # use first row label for style decision (sick vs annual)
        first_label = rows[0]["shift"] if rows else ""
        title, icon, border, bg, text_color, count_bg = shift_style(g, first_label)

        # open only one group if requested
        open_attr = " open" if (open_group and g == open_group) else ""

        emp_rows_html = []
        for i, x in enumerate(rows):
            alt = " empRowAlt" if i % 2 == 1 else ""
            emp_rows_html.append(
                f"""<div class="empRow{alt}">
       <span class="empName">{x["name"]}</span>
       <span class="empStatus" style="color:{text_color};">{x["shift"]}</span>
     </div>"""
            )

        shift_blocks.append(
            f"""
    <details class="shiftCard" style="border:1px solid {border}; background:{bg};"{open_attr}>
      <summary class="shiftSummary" style="background:{bg}; border-bottom:1px solid {border.replace('44','33')};">
        <span class="shiftIcon">{icon}</span>
        <span class="shiftLabel" style="color:{text_color};">{title}</span>
        <span class="shiftCount" style="background:{count_bg}; color:{text_color};">{len(rows)}</span>
      </summary>
      <div class="shiftBody">
        {''.join(emp_rows_html)}
      </div>
    </details>
            """
        )

    if not shift_blocks:
        shift_blocks_html = '<div class="shiftStack"><div class="footer" style="margin:0; padding:14px 0;">No data for today</div></div>'
    else:
        shift_blocks_html = f'<div class="shiftStack">{"".join(shift_blocks)}</div>'

    return f"""
    <div class="deptCard">
      <div style="height:5px; background:linear-gradient(to right, {dept_color}, {dept_color}cc);"></div>

      <div class="deptHead" style="border-bottom:2px solid {dept_color}18;">
        <div class="deptIcon" style="background:{dept_color}15; color:{dept_color};">
          {SVG_ICON}
        </div>
        <div class="deptTitle">{dept_name}</div>
        <div class="deptBadge" style="background:{dept_color}12; color:{dept_color}; border:1px solid {dept_color}28;">
          <span style="font-size:10px;opacity:.7;display:block;margin-bottom:1px;text-transform:uppercase;letter-spacing:.5px;">Total</span>
          <span style="font-size:17px;font-weight:900;">{total}</span>
        </div>
      </div>

      {shift_blocks_html}
    </div>
    """

def page_shell_html(date_label: str, employees_total: int, departments_total: int, dept_cards_html: str, cta_url: str, sent_time: str):
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="x-apple-disable-message-reformatting">
  <title>Duty Roster</title>
  <style>
{CSS}
  </style>
</head>
<body>
<div class="wrap">

  <!-- ════ HEADER ════ -->
  <div class="header">
    <h1>📋 Duty Roster</h1>
    <div class="dateTag">📅 {date_label}</div>
  </div>

  <!-- ════ SUMMARY CHIPS ════ -->
  <div class="summaryBar">
    <div class="summaryChip">
      <div class="chipVal">{employees_total}</div>
      <div class="chipLabel">Employees</div>
    </div>
    <div class="summaryChip">
      <div class="chipVal" style="color:#059669;">{departments_total}</div>
      <div class="chipLabel">Departments</div>
    </div>
  </div>

  <!-- ════ DEPARTMENT CARDS ════ -->
  {dept_cards_html}

  <!-- ════ CTA ════ -->
  <div class="btnWrap">
    <a class="btn" href="{cta_url}">📋 View Full Duty Roster</a>
  </div>

  <!-- ════ FOOTER ════ -->
  <div class="footer">
    Sent at <strong>{sent_time}</strong>
     &nbsp;·&nbsp; Total: <strong>{employees_total} employees</strong>
  </div>

</div>
</body>
</html>"""


def build_pretty_email_html(active_group: str, now: datetime, rows_by_dept: list, pages_base: str) -> str:
    """
    Email-safe HTML (tables + inline styles) with department header colors like the site.
    rows_by_dept = [{"dept": str, "rows": [{"name": str, "shift": str}, ...]}, ...]
    """
    # Date label (robust for runners)
    try:
        date_label = now.strftime("%-d %B %Y")
    except Exception:
        date_label = now.strftime("%d %B %Y")

    sent_time = now.strftime("%H:%M")

    # Shift theme (for status color)
    def shift_theme(g: str):
        if g == "صباح":
            return ("#fef3c7", "#f59e0b55", "#92400e")
        if g == "ظهر":
            return ("#ffedd5", "#f9731655", "#9a3412")
        if g == "ليل":
            return ("#ede9fe", "#8b5cf655", "#5b21b6")
        return ("#e0e7ff", "#6366f155", "#3730a3")

    bg, border, textc = shift_theme(active_group)

    dept_blocks = []
    total_now = 0
    depts_now = 0

    for item in rows_by_dept:
        dept = item.get("dept", "")
        rows = item.get("rows", []) or []
        if not rows:
            continue

        depts_now += 1
        total_now += len(rows)

        trs = []
        for i, r in enumerate(rows):
            alt_bg = "#f8fafc" if i % 2 == 1 else "#ffffff"
            trs.append(f"""
              <tr>
                <td style="padding:10px 12px;border-top:1px solid #eef2f7;background:{alt_bg};font-weight:700;color:#0f172a;">
                  {r["name"]}
                </td>
                <td style="padding:10px 12px;border-top:1px solid #eef2f7;background:{alt_bg};white-space:nowrap;font-weight:700;color:{textc};">
                  {r["shift"]}
                </td>
              </tr>
            """)

        dept_color = DEPT_EMAIL_COLORS.get(dept, "#1e40af")

        dept_blocks.append(f"""
          <table role="presentation" width="100%" cellpadding="0" cellspacing="0"
                 style="margin-top:16px;border:1px solid #e6e6e6;border-radius:16px;overflow:hidden;background:#ffffff;">
            <tr>
              <td style="height:6px;background:{dept_color};font-size:0;line-height:0;">&nbsp;</td>
            </tr>
            <tr>
              <td style="padding:12px 14px;border-bottom:1px solid #eef2f7;">
                <table role="presentation" width="100%" cellpadding="0" cellspacing="0">
                  <tr>
                    <td style="font-size:16px;font-weight:900;color:{dept_color};">
                      {dept}
                    </td>
                    <td align="right">
                      <span style="
                        display:inline-block;
                        padding:6px 12px;
                        border-radius:12px;
                        font-size:13px;
                        font-weight:900;
                        color:{dept_color};
                        background:{dept_color}22;
                        border:1px solid {dept_color}55;
                      ">
                        TOTAL {len(rows)}
                      </span>
                    </td>
                  </tr>
                </table>
              </td>
            </tr>
            <tr>
              <td>
                <table role="presentation" width="100%" cellpadding="0" cellspacing="0">
                  <tr style="background:#f6f7f9;">
                    <th align="left" style="padding:10px 14px;border-bottom:1px solid #eef2f7;color:#334155;font-size:12px;letter-spacing:.4px;text-transform:uppercase;">
                      Employee
                    </th>
                    <th align="left" style="padding:10px 14px;border-bottom:1px solid #eef2f7;color:#334155;font-size:12px;letter-spacing:.4px;text-transform:uppercase;">
                      Status
                    </th>
                  </tr>
                  {''.join(trs)}
                </table>
              </td>
            </tr>
          </table>
        """)

    dept_html = "\n".join(dept_blocks) if dept_blocks else f"""
      <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="margin-top:14px;">
        <tr>
          <td style="padding:14px;border-radius:14px;border:1px dashed rgba(15,23,42,.18);background:#ffffff;">
            <div style="font-weight:900;color:#334155;">No staff for current shift.</div>
            <div style="margin-top:6px;color:#64748b;font-size:13px;">Open the website for full details.</div>
          </td>
        </tr>
      </table>
    """

    pages_base = (pages_base or "").rstrip("/")

    return f"""<!doctype html>
<html>
  <body style="margin:0;padding:0;background:#eef1f7;font-family:Segoe UI,Arial,Helvetica,sans-serif;color:#0f172a;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#eef1f7;">
      <tr>
        <td align="center" style="padding:16px 10px;">
          <table role="presentation" width="680" cellpadding="0" cellspacing="0" style="max-width:680px;width:100%;">
            <tr>
              <td style="border-radius:20px;overflow:hidden;box-shadow:0 8px 28px rgba(30,64,175,.18);">

                <div style="background:linear-gradient(135deg,#1e40af 0%,#1976d2 50%,#0ea5e9 100%);padding:22px 18px;color:#fff;text-align:center;">
                  <div style="font-size:22px;font-weight:900;letter-spacing:-.2px;">📋 Duty Roster</div>
                  <div style="margin-top:8px;display:inline-block;background:rgba(255,255,255,.18);padding:6px 16px;border-radius:30px;font-size:13px;font-weight:700;">
                    📅 {date_label}
                  </div>
                </div>

                <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#ffffff;">
                  <tr>
                    <td style="padding:16px 16px 10px 16px;">

                      <div style="margin:0 auto 12px auto;display:inline-block;padding:10px 14px;border-radius:14px;background:{bg};border:1px solid {border};color:{textc};font-weight:900;">
                        Current shift: {active_group}
                      </div>

                      <table role="presentation" cellpadding="0" cellspacing="0" style="width:100%;margin-top:6px;">
                        <tr>
                          <td style="width:50%;padding-right:6px;">
                            <div style="border:1px solid rgba(15,23,42,.10);border-radius:14px;padding:10px 12px;text-align:center;background:#fff;">
                              <div style="font-size:22px;font-weight:900;color:#1e40af;">{total_now}</div>
                              <div style="font-size:11px;font-weight:700;color:#64748b;letter-spacing:.5px;text-transform:uppercase;">Now</div>
                            </div>
                          </td>
                          <td style="width:50%;padding-left:6px;">
                            <div style="border:1px solid rgba(15,23,42,.10);border-radius:14px;padding:10px 12px;text-align:center;background:#fff;">
                              <div style="font-size:22px;font-weight:900;color:#059669;">{depts_now}</div>
                              <div style="font-size:11px;font-weight:700;color:#64748b;letter-spacing:.5px;text-transform:uppercase;">Departments</div>
                            </div>
                          </td>
                        </tr>
                      </table>

                      {dept_html}

                      <div style="text-align:center;margin-top:16px;">
                        <a href="{pages_base}/now/" style="display:inline-block;padding:12px 18px;border-radius:16px;background:linear-gradient(135deg,#1e40af,#1976d2);color:#fff;text-decoration:none;font-weight:900;">
                          Open Now Page
                        </a>
                        <span style="display:inline-block;width:10px;"></span>
                        <a href="{pages_base}/" style="display:inline-block;padding:12px 18px;border-radius:16px;background:#0ea5e9;color:#fff;text-decoration:none;font-weight:900;">
                          Open Full Page
                        </a>
                      </div>

                      <div style="margin-top:14px;text-align:center;color:#94a3b8;font-size:12px;line-height:1.9;">
                        Sent at <strong style="color:#64748b;">{sent_time}</strong> · GitHub Actions
                      </div>

                    </td>
                  </tr>
                </table>

              </td>
            </tr>
          </table>
        </td>
      </tr>
    </table>
  </body>
</html>"""





# =========================
# Email
# =========================
def send_email(subject: str, html: str):
    if not (SMTP_HOST and SMTP_USER and SMTP_PASS and MAIL_FROM and MAIL_TO):
        return
    msg = MIMEText(html, "html", "utf-8")
    msg["Subject"] = subject
    msg["From"] = MAIL_FROM
    msg["To"] = MAIL_TO

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as s:
        s.starttls()
        s.login(SMTP_USER, SMTP_PASS)
        s.sendmail(MAIL_FROM, [x.strip() for x in MAIL_TO.split(",") if x.strip()], msg.as_string())


# =========================
# Main
# =========================
def main():
    if not EXCEL_URL:
        raise RuntimeError("EXCEL_URL missing")

    now = datetime.now(TZ)
    # Sun=0..Sat=6
    today_dow = (now.weekday() + 1) % 7
    today_day = now.day

    active_group = current_shift_key(now)  # "صباح" / "ظهر" / "ليل"
    pages_base = (PAGES_BASE_URL or infer_pages_base_url()).rstrip("/")

    data = download_excel(EXCEL_URL)
    wb = load_workbook(BytesIO(data), data_only=True)

    dept_cards_all = []
    dept_cards_now = []
    rows_by_dept = []  # for email (NOW only)
    employees_total_all = 0
    employees_total_now = 0
    depts_count = 0

    for idx, (sheet_name, dept_name) in enumerate(DEPARTMENTS):
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        days_row, date_row = find_days_and_dates_rows(ws)
        day_col = find_day_col(ws, days_row, date_row, today_dow, today_day)

        if not (days_row and date_row and day_col):
            # skip if sheet layout unexpected
            continue

        start_row = date_row + 1
        emp_col = find_employee_col(ws, start_row=start_row)
        if not emp_col:
            continue

        buckets = {k: [] for k in GROUP_ORDER}
        buckets_now = {k: [] for k in GROUP_ORDER}

        for r in range(start_row, ws.max_row + 1):
            name = norm(ws.cell(row=r, column=emp_col).value)
            if not looks_like_employee_name(name):
                continue

            raw = norm(ws.cell(row=r, column=day_col).value)
            if not looks_like_shift_code(raw):
                continue

            label, grp = map_shift(raw)
            buckets.setdefault(grp, []).append({"name": name, "shift": label})

            if grp == active_group:
                buckets_now.setdefault(grp, []).append({"name": name, "shift": label})

        # Collect NOW rows for email (current shift only)
        now_rows = buckets_now.get(active_group, [])
        rows_by_dept.append({"dept": dept_name, "rows": now_rows})

        dept_color = DEPT_COLORS[idx % len(DEPT_COLORS)]
        open_group_full = active_group if AUTO_OPEN_ACTIVE_SHIFT_IN_FULL else None
        card_all = dept_card_html(dept_name, dept_color, buckets, open_group=open_group_full)
        dept_cards_all.append(card_all)

        # For NOW page: open the active shift group by default
        card_now = dept_card_html(dept_name, dept_color, buckets_now, open_group=active_group)
        dept_cards_now.append(card_now)

        employees_total_all += sum(len(buckets.get(g, [])) for g in GROUP_ORDER)
        employees_total_now += sum(len(buckets_now.get(g, [])) for g in GROUP_ORDER)

        depts_count += 1

    # pages
    os.makedirs("docs", exist_ok=True)
    os.makedirs("docs/now", exist_ok=True)

    date_label = now.strftime("%-d %B %Y") if hasattr(now, "strftime") else now.strftime("%d %B %Y")
    # Windows runners sometimes don't support %-d; safe fallback
    try:
        date_label = now.strftime("%-d %B %Y")
    except Exception:
        date_label = now.strftime("%d %B %Y")

    sent_time = now.strftime("%H:%M")

    full_url = f"{pages_base}/"
    now_url = f"{pages_base}/now/"

    html_full = page_shell_html(
        date_label=date_label,
        employees_total=employees_total_all,
        departments_total=depts_count,
        dept_cards_html="\n".join(dept_cards_all),
        cta_url=now_url,   # button on full page goes to NOW page
        sent_time=sent_time,
    )
    html_now = page_shell_html(
        date_label=date_label,
        employees_total=employees_total_now,
        departments_total=depts_count,
        dept_cards_html="\n".join(dept_cards_now),
        cta_url=full_url,  # button on now page goes to FULL page
        sent_time=sent_time,
    )

    with open("docs/index.html", "w", encoding="utf-8") as f:
        f.write(html_full)

    with open("docs/now/index.html", "w", encoding="utf-8") as f:
        f.write(html_now)

    # Email: send a dedicated email-safe template (better rendering in Gmail/Outlook)
    subject = f"Duty Roster — {active_group} — {now.strftime('%Y-%m-%d')}"
    email_html = build_pretty_email_html(active_group, now, rows_by_dept, pages_base)
    send_email(subject, email_html)


if __name__ == "__main__":
    main()